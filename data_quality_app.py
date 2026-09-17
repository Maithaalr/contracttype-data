import streamlit as st
import pandas as pd
from io import BytesIO

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)


# =========================================================
# إعداد الصفحة
# =========================================================

st.set_page_config(
    page_title="نظام تحليل العقود",
    page_icon="📊",
    layout="wide"
)


# =========================================================
# تنسيق الصفحة RTL
# =========================================================

st.markdown("""
<style>

    html, body, [class*="css"] {
        direction: rtl;
        text-align: right;
    }

    .main {
        direction: rtl;
    }

    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }

    div[data-testid="stMetric"] {
        background-color: #f8f9fa;
        border: 1px solid #e9ecef;
        padding: 18px;
        border-radius: 12px;
        text-align: center;
    }

    div[data-testid="stMetricLabel"] {
        justify-content: center;
    }

    div[data-testid="stMetricValue"] {
        text-align: center;
    }

    .main-title {
        font-size: 32px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        color: #6c757d;
        font-size: 16px;
        margin-bottom: 25px;
    }

    .section-title {
        font-size: 21px;
        font-weight: 700;
        margin-top: 25px;
        margin-bottom: 15px;
    }

</style>
""", unsafe_allow_html=True)


# =========================================================
# دوال مساعدة
# =========================================================

def clean_column_names(df):

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
        .str.replace("\n", " ", regex=False)
    )

    return df


def find_column(df, possible_names):

    normalized = {
        str(col).strip().replace(" ", ""): col
        for col in df.columns
    }

    for name in possible_names:

        key = name.strip().replace(" ", "")

        if key in normalized:
            return normalized[key]

    return None


def convert_numeric(df, column):

    if column and column in df.columns:

        df[column] = (
            df[column]
            .astype(str)
            .str.replace(",", "", regex=False)
            .str.replace("AED", "", regex=False)
            .str.replace("درهم", "", regex=False)
            .str.strip()
        )

        df[column] = pd.to_numeric(
            df[column],
            errors="coerce"
        ).fillna(0)

    return df


def format_money(value):

    try:
        return f"{value:,.2f}"

    except:
        return "0.00"


# =========================================================
# دالة إنشاء ملف Excel
# =========================================================

def create_excel(
    filtered_df,
    contract_type,
    total_contracts,
    total_cost,
    total_basic,
    total_supplementary,
    top_entity,
    top_category,
    category_counts,
    entity_counts,
    nature_analysis,
    total_nature_all,
    total_nature_eligible_all
):

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        # =================================================
        # SHEET 1 - بيانات العقود المختارة
        # =================================================

        filtered_df.to_excel(
            writer,
            sheet_name="العقود",
            index=False
        )


        # =================================================
        # SHEET 2 - الإحصائيات
        # =================================================

        summary = pd.DataFrame({

            "المؤشر": [
                "نوع العقد المختار",
                "إجمالي عدد العقود",
                "أكثر جهة حكومية",
                "أكثر فئة وظيفية",
                "إجمالي الراتب الأساسي",
                "إجمالي التكميلي",
                "إجمالي التكلفة الشهرية"
            ],

            "القيمة": [
                contract_type,
                total_contracts,
                top_entity,
                top_category,
                total_basic,
                total_supplementary,
                total_cost
            ]

        })


        summary.to_excel(
            writer,
            sheet_name="الإحصائيات",
            index=False,
            startrow=0
        )


        # =================================================
        # توزيع الموظفين حسب الفئة الوظيفية
        # =================================================

        category_start_row = len(summary) + 3


        if not category_counts.empty:

            category_export = (
                category_counts
                .reset_index()
            )

            category_export.columns = [
                "الفئة الوظيفية",
                "عدد الموظفين"
            ]


            category_total = pd.DataFrame({

                "الفئة الوظيفية": [
                    "الإجمالي"
                ],

                "عدد الموظفين": [
                    category_export[
                        "عدد الموظفين"
                    ].sum()
                ]

            })


            category_export = pd.concat(
                [
                    category_export,
                    category_total
                ],
                ignore_index=True
            )


            category_export.to_excel(
                writer,
                sheet_name="الإحصائيات",
                index=False,
                startrow=category_start_row
            )

        else:

            category_export = pd.DataFrame()


        # =================================================
        # عدد الموظفين حسب الجهة الحكومية
        # =================================================

        entity_start_row = (
            category_start_row
            + len(category_export)
            + 4
        )


        if not entity_counts.empty:

            entity_export = (
                entity_counts
                .reset_index()
            )

            entity_export.columns = [
                "الجهة الحكومية",
                "عدد الموظفين"
            ]


            entity_total = pd.DataFrame({

                "الجهة الحكومية": [
                    "الإجمالي"
                ],

                "عدد الموظفين": [
                    entity_export[
                        "عدد الموظفين"
                    ].sum()
                ]

            })


            entity_export = pd.concat(
                [
                    entity_export,
                    entity_total
                ],
                ignore_index=True
            )


            entity_export.to_excel(
                writer,
                sheet_name="الإحصائيات",
                index=False,
                startrow=entity_start_row
            )

        else:

            entity_export = pd.DataFrame()


        # =================================================
        # تحليل بدل طبيعة العمل - جميع أنواع العقود
        # =================================================

        nature_start_row = (
            entity_start_row
            + len(entity_export)
            + 5
        )


        # ملخص بدل طبيعة العمل
        nature_summary = pd.DataFrame({

            "المؤشر": [
                "عدد الموظفين المستحقين لبدل طبيعة العمل",
                "إجمالي قيمة بدل طبيعة العمل"
            ],

            "القيمة": [
                total_nature_eligible_all,
                total_nature_all
            ]

        })


        nature_summary.to_excel(
            writer,
            sheet_name="الإحصائيات",
            index=False,
            startrow=nature_start_row
        )


        nature_table_start = (
            nature_start_row
            + len(nature_summary)
            + 3
        )


        if not nature_analysis.empty:

            nature_export = (
                nature_analysis.copy()
            )


            nature_total = pd.DataFrame({

                "الجهة الحكومية": [
                    "الإجمالي"
                ],

                "الفئة الوظيفية": [
                    ""
                ],

                "قيمة البدل": [
                    ""
                ],

                "عدد الموظفين المستحقين": [
                    nature_export[
                        "عدد الموظفين المستحقين"
                    ].sum()
                ],

                "إجمالي قيمة البدل": [
                    nature_export[
                        "إجمالي قيمة البدل"
                    ].sum()
                ]

            })


            nature_export = pd.concat(
                [
                    nature_export,
                    nature_total
                ],
                ignore_index=True
            )


            nature_export.to_excel(
                writer,
                sheet_name="الإحصائيات",
                index=False,
                startrow=nature_table_start
            )


        # =================================================
        # تنسيق Excel
        # =================================================

        workbook = writer.book


        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )


        header_font = Font(
            bold=True
        )


        total_fill = PatternFill(
            fill_type="solid",
            fgColor="E2F0D9"
        )


        thin_border = Border(

            left=Side(
                style="thin",
                color="D9D9D9"
            ),

            right=Side(
                style="thin",
                color="D9D9D9"
            ),

            top=Side(
                style="thin",
                color="D9D9D9"
            ),

            bottom=Side(
                style="thin",
                color="D9D9D9"
            )

        )


        for sheet_name in workbook.sheetnames:

            worksheet = workbook[sheet_name]

            # اتجاه عربي
            worksheet.sheet_view.rightToLeft = True

            # تثبيت أول صف
            worksheet.freeze_panes = "A2"


            # محاذاة وحدود
            for row in worksheet.iter_rows():

                for cell in row:

                    cell.alignment = Alignment(
                        horizontal="right",
                        vertical="center"
                    )

                    if cell.value is not None:

                        cell.border = thin_border


            # تنسيق صفوف العناوين
            header_names = [
                "المؤشر",
                "الفئة الوظيفية",
                "الجهة الحكومية"
            ]


            for row in worksheet.iter_rows():

                first_value = row[0].value

                if first_value in header_names:

                    for cell in row:

                        if cell.value is not None:

                            cell.fill = header_fill
                            cell.font = header_font


            # تنسيق صفوف الإجمالي
            for row in worksheet.iter_rows():

                if any(
                    cell.value == "الإجمالي"
                    for cell in row
                ):

                    for cell in row:

                        if cell.value is not None:

                            cell.fill = total_fill
                            cell.font = Font(
                                bold=True
                            )


            # عرض الأعمدة
            for column_cells in worksheet.columns:

                max_length = 0

                column_letter = (
                    column_cells[0]
                    .column_letter
                )


                for cell in column_cells:

                    try:

                        if cell.value is not None:

                            length = len(
                                str(cell.value)
                            )

                            if length > max_length:

                                max_length = length

                    except:

                        pass


                adjusted_width = min(
                    max_length + 5,
                    45
                )


                worksheet.column_dimensions[
                    column_letter
                ].width = adjusted_width


            # ارتفاع الصفوف
            for row_number in range(
                1,
                worksheet.max_row + 1
            ):

                worksheet.row_dimensions[
                    row_number
                ].height = 22


    output.seek(0)

    return output


# =========================================================
# عنوان النظام
# =========================================================

st.markdown(
    '<div class="main-title">'
    'نظام تحليل العقود'
    '</div>',
    unsafe_allow_html=True
)


st.markdown(
    '<div class="subtitle">'
    'تحليل أنواع العقود والتكاليف والإحصائيات بصورة تلقائية'
    '</div>',
    unsafe_allow_html=True
)


# =========================================================
# رفع ملف Excel
# =========================================================

uploaded_file = st.file_uploader(
    "ارفع ملف بيانات العقود",
    type=["xlsx", "xls"]
)


if uploaded_file is None:

    st.info(
        "يرجى رفع ملف Excel الذي يحتوي على بيانات العقود للبدء."
    )

    st.stop()


# =========================================================
# قراءة ملف Excel
# =========================================================

try:

    excel_file = pd.ExcelFile(
        uploaded_file
    )

    sheet_names = (
        excel_file.sheet_names
    )


    if len(sheet_names) > 1:

        selected_sheet = st.selectbox(
            "اختر شيت البيانات",
            sheet_names
        )

    else:

        selected_sheet = (
            sheet_names[0]
        )


    df = pd.read_excel(
        uploaded_file,
        sheet_name=selected_sheet
    )


except Exception as e:

    st.error(
        f"حدث خطأ أثناء قراءة الملف: {e}"
    )

    st.stop()


# =========================================================
# تنظيف البيانات
# =========================================================

df = clean_column_names(df)

df = df.dropna(
    how="all"
)


# =========================================================
# اكتشاف الأعمدة تلقائياً
# =========================================================

contract_type_col = find_column(
    df,
    [
        "نوع العقد",
        "نوع عقد",
        "العقد",
        "Contract Type"
    ]
)


entity_col = find_column(
    df,
    [
        "الدائرة",
        "الجهة",
        "الجهة الحكومية",
        "اسم الجهة",
        "اسم الدائرة",
        "الدائرة الحكومية",
        "Entity"
    ]
)


category_col = find_column(
    df,
    [
        "الفئة الوظيفية",
        "الفئة",
        "Job Category"
    ]
)


basic_salary_col = find_column(
    df,
    [
        "الراتب الأساسي",
        "الراتب الاساسي",
        "أساسي",
        "Basic Salary"
    ]
)


supplementary_col = find_column(
    df,
    [
        "التكميلي",
        "الراتب التكميلي",
        "Supplementary"
    ]
)


nature_allowance_col = find_column(
    df,
    [
        "بدل طبيعة العمل",
        "بدل طبيعة عمل",
        "Nature Allowance"
    ]
)


monthly_cost_col = find_column(
    df,
    [
        "التكلفة الشهرية",
        "إجمالي التكلفة الشهرية",
        "اجمالي التكلفة الشهرية",
        "التكلفة",
        "Monthly Cost"
    ]
)


# =========================================================
# التأكد من وجود نوع العقد
# =========================================================

if contract_type_col is None:

    st.error(
        """
        لم أجد عمود **نوع العقد** في الملف.

        يجب أن يحتوي الملف على عمود باسم **نوع العقد**.
        """
    )

    st.write(
        "الأعمدة الموجودة في الملف:"
    )

    st.write(
        list(df.columns)
    )

    st.stop()


# =========================================================
# تنظيف نوع العقد
# =========================================================

df[contract_type_col] = (
    df[contract_type_col]
    .astype(str)
    .str.strip()
)


df = df[
    ~df[contract_type_col].isin(
        [
            "",
            "nan",
            "None"
        ]
    )
]


# =========================================================
# تحويل الأعمدة المالية إلى أرقام
# =========================================================

for col in [
    basic_salary_col,
    supplementary_col,
    nature_allowance_col,
    monthly_cost_col
]:

    df = convert_numeric(
        df,
        col
    )


# =========================================================
# اختيار نوع العقد
# =========================================================

st.markdown(
    '<div class="section-title">'
    'اختيار نوع العقد'
    '</div>',
    unsafe_allow_html=True
)


contract_types = sorted(
    df[
        contract_type_col
    ]
    .dropna()
    .unique()
    .tolist()
)


selected_contract = st.selectbox(
    "نوع العقد",
    contract_types
)


# =========================================================
# فلترة البيانات حسب نوع العقد
#
# هذا الفلتر يستخدم لكل شيء
# ما عدا تحليل بدل طبيعة العمل
# =========================================================

filtered_df = df[
    df[contract_type_col]
    == selected_contract
].copy()


# =========================================================
# حساب التكلفة الشهرية إذا لم تكن موجودة
# =========================================================

if monthly_cost_col is None:

    available_salary_columns = [

        col

        for col in [
            basic_salary_col,
            supplementary_col,
            nature_allowance_col
        ]

        if col is not None
    ]


    if available_salary_columns:

        filtered_df[
            "التكلفة الشهرية المحسوبة"
        ] = (

            filtered_df[
                available_salary_columns
            ].sum(axis=1)

        )

        monthly_cost_col = (
            "التكلفة الشهرية المحسوبة"
        )


# =========================================================
# إحصائيات نوع العقد المختار
# =========================================================

total_contracts = len(
    filtered_df
)


# الراتب الأساسي
if basic_salary_col:

    total_basic = (
        filtered_df[
            basic_salary_col
        ].sum()
    )

else:

    total_basic = 0


# التكميلي
if supplementary_col:

    total_supplementary = (
        filtered_df[
            supplementary_col
        ].sum()
    )

else:

    total_supplementary = 0


# التكلفة الشهرية
if monthly_cost_col:

    total_cost = (
        filtered_df[
            monthly_cost_col
        ].sum()
    )

else:

    total_cost = (
        total_basic
        + total_supplementary
    )


# =========================================================
# أكثر جهة حكومية لنوع العقد المختار
# =========================================================

if entity_col and not filtered_df.empty:

    entity_data = (
        filtered_df[
            entity_col
        ]
        .dropna()
        .astype(str)
        .str.strip()
    )


    entity_data = entity_data[
        (entity_data != "")
        & (entity_data != "nan")
    ]


    if not entity_data.empty:

        top_entity = (
            entity_data
            .value_counts()
            .idxmax()
        )

    else:

        top_entity = "غير متوفر"

else:

    top_entity = "غير متوفر"


# =========================================================
# عدد الموظفين حسب الجهة الحكومية
# لنوع العقد المختار فقط
# =========================================================

if entity_col and not filtered_df.empty:

    entity_counts = (
        filtered_df[
            entity_col
        ]
        .dropna()
        .astype(str)
        .str.strip()
    )


    entity_counts = entity_counts[
        (entity_counts != "")
        & (entity_counts != "nan")
    ]


    entity_counts = (
        entity_counts
        .value_counts()
    )

else:

    entity_counts = (
        pd.Series(dtype=int)
    )


# =========================================================
# الفئات الوظيفية لنوع العقد المختار
# =========================================================

if category_col and not filtered_df.empty:

    category_data = (
        filtered_df[
            category_col
        ]
        .dropna()
        .astype(str)
        .str.strip()
    )


    category_data = category_data[
        (category_data != "")
        & (category_data != "nan")
    ]


    if not category_data.empty:

        category_counts = (
            category_data
            .value_counts()
        )


        top_category = (
            category_counts
            .idxmax()
        )

    else:

        category_counts = (
            pd.Series(dtype=int)
        )

        top_category = (
            "غير متوفر"
        )

else:

    category_counts = (
        pd.Series(dtype=int)
    )

    top_category = (
        "غير متوفر"
    )


# =========================================================
# تحليل بدل طبيعة العمل
#
# مهم جداً:
# هذا القسم لا يستخدم filtered_df
# وإنما يستخدم df بالكامل
#
# لذلك يشمل جميع أنواع العقود
# ولا يتأثر باختيار نوع العقد
# =========================================================

nature_analysis = pd.DataFrame()

total_nature_all = 0

total_nature_eligible_all = 0


if (
    nature_allowance_col
    and entity_col
    and category_col
    and not df.empty
):

    # =====================================================
    # جميع الموظفين الذين بدل طبيعة العمل لديهم أكبر من صفر
    # من جميع أنواع العقود
    # =====================================================

    nature_df = df[
        df[
            nature_allowance_col
        ] > 0
    ].copy()


    # =====================================================
    # إجمالي عدد المستحقين
    # =====================================================

    total_nature_eligible_all = (
        len(nature_df)
    )


    # =====================================================
    # إجمالي قيمة البدل
    # =====================================================

    total_nature_all = (
        nature_df[
            nature_allowance_col
        ].sum()
    )


    # =====================================================
    # تنظيف اسم الجهة
    # =====================================================

    nature_df[entity_col] = (
        nature_df[
            entity_col
        ]
        .astype(str)
        .str.strip()
    )


    # =====================================================
    # تنظيف الفئة الوظيفية
    # =====================================================

    nature_df[category_col] = (
        nature_df[
            category_col
        ]
        .astype(str)
        .str.strip()
    )


    # =====================================================
    # حذف الجهة أو الفئة الفارغة
    # =====================================================

    nature_df = nature_df[

        (nature_df[entity_col] != "")
        & (nature_df[entity_col] != "nan")

        &

        (nature_df[category_col] != "")
        & (nature_df[category_col] != "nan")

    ]


    # =====================================================
    # التجميع
    # الجهة + الفئة + قيمة البدل
    # =====================================================

    if not nature_df.empty:

        nature_analysis = (

            nature_df

            .groupby(
                [
                    entity_col,
                    category_col,
                    nature_allowance_col
                ],
                dropna=False
            )

            .size()

            .reset_index(
                name="عدد الموظفين المستحقين"
            )

        )


        # =================================================
        # إجمالي قيمة البدل لكل مجموعة
        # =================================================

        nature_analysis[
            "إجمالي قيمة البدل"
        ] = (

            nature_analysis[
                nature_allowance_col
            ]

            *

            nature_analysis[
                "عدد الموظفين المستحقين"
            ]

        )


        # =================================================
        # أسماء الأعمدة
        # =================================================

        nature_analysis.columns = [

            "الجهة الحكومية",

            "الفئة الوظيفية",

            "قيمة البدل",

            "عدد الموظفين المستحقين",

            "إجمالي قيمة البدل"

        ]


        # =================================================
        # ترتيب الجدول
        # =================================================

        nature_analysis = (

            nature_analysis

            .sort_values(
                by=[
                    "الجهة الحكومية",
                    "الفئة الوظيفية",
                    "قيمة البدل"
                ]
            )

            .reset_index(
                drop=True
            )

        )


# =========================================================
# DASHBOARD
# =========================================================

st.markdown("---")


st.markdown(
    f'<div class="section-title">'
    f'ملخص {selected_contract}'
    f'</div>',
    unsafe_allow_html=True
)


# =========================================================
# البطاقات الرئيسية
# =========================================================

col1, col2, col3, col4 = (
    st.columns(4)
)


with col1:

    st.metric(
        "إجمالي العقود",
        f"{total_contracts:,}"
    )


with col2:

    st.metric(
        "إجمالي التكلفة الشهرية",
        f"{format_money(total_cost)} د.إ"
    )


with col3:

    st.metric(
        "أكثر جهة حكومية",
        top_entity
    )


with col4:

    st.metric(
        "أكثر فئة وظيفية",
        top_category
    )


# =========================================================
# تفاصيل التكلفة للعقد المختار
# =========================================================

st.markdown(
    '<div class="section-title">'
    'تفاصيل التكلفة الشهرية'
    '</div>',
    unsafe_allow_html=True
)


cost1, cost2 = (
    st.columns(2)
)


with cost1:

    st.metric(
        "إجمالي الراتب الأساسي",
        f"{format_money(total_basic)} د.إ"
    )


with cost2:

    st.metric(
        "إجمالي التكميلي",
        f"{format_money(total_supplementary)} د.إ"
    )


# =========================================================
# توزيع الموظفين حسب الفئة الوظيفية
# =========================================================

if not category_counts.empty:

    st.markdown(
        '<div class="section-title">'
        'توزيع الموظفين حسب الفئة الوظيفية'
        '</div>',
        unsafe_allow_html=True
    )


    category_df = (
        category_counts
        .reset_index()
    )


    category_df.columns = [
        "الفئة الوظيفية",
        "عدد الموظفين"
    ]


    category_total = pd.DataFrame({

        "الفئة الوظيفية": [
            "الإجمالي"
        ],

        "عدد الموظفين": [
            category_df[
                "عدد الموظفين"
            ].sum()
        ]

    })


    category_df_display = pd.concat(
        [
            category_df,
            category_total
        ],
        ignore_index=True
    )


    category_chart_col, category_table_col = (
        st.columns([2, 1])
    )


    with category_chart_col:

        st.bar_chart(
            category_df.set_index(
                "الفئة الوظيفية"
            )
        )


    with category_table_col:

        st.dataframe(
            category_df_display,
            use_container_width=True,
            hide_index=True
        )


# =========================================================
# عدد الموظفين حسب الجهة الحكومية
# =========================================================

if not entity_counts.empty:

    st.markdown(
        '<div class="section-title">'
        'عدد الموظفين حسب الجهة الحكومية'
        '</div>',
        unsafe_allow_html=True
    )


    entity_df = (
        entity_counts
        .reset_index()
    )


    entity_df.columns = [
        "الجهة الحكومية",
        "عدد الموظفين"
    ]


    entity_total = pd.DataFrame({

        "الجهة الحكومية": [
            "الإجمالي"
        ],

        "عدد الموظفين": [
            entity_df[
                "عدد الموظفين"
            ].sum()
        ]

    })


    entity_df_display = pd.concat(
        [
            entity_df,
            entity_total
        ],
        ignore_index=True
    )


    entity_chart_col, entity_table_col = (
        st.columns([2, 1])
    )


    with entity_chart_col:

        st.bar_chart(
            entity_df.set_index(
                "الجهة الحكومية"
            )
        )


    with entity_table_col:

        st.dataframe(
            entity_df_display,
            use_container_width=True,
            hide_index=True
        )


# =========================================================
# تحليل بدل طبيعة العمل
#
# هذا القسم لجميع أنواع العقود
# =========================================================

st.markdown("---")


st.markdown(
    '<div class="section-title">'
    'تحليل بدل طبيعة العمل'
    '</div>',
    unsafe_allow_html=True
)


st.caption(
    "يشمل هذا التحليل جميع أنواع العقود ولا يتأثر بفلتر نوع العقد."
)


if nature_allowance_col is None:

    st.warning(
        "لم يتم العثور على عمود بدل طبيعة العمل في البيانات."
    )


elif nature_analysis.empty:

    st.info(
        "لا يوجد موظفون مستحقون لبدل طبيعة العمل في البيانات."
    )


else:

    # =====================================================
    # بطاقات بدل طبيعة العمل
    # =====================================================

    nature1, nature2 = (
        st.columns(2)
    )


    with nature1:

        st.metric(
            "إجمالي عدد الموظفين المستحقين",
            f"{total_nature_eligible_all:,}"
        )


    with nature2:

        st.metric(
            "إجمالي قيمة بدل طبيعة العمل",
            f"{format_money(total_nature_all)} د.إ"
        )


    # =====================================================
    # الجدول التفصيلي
    # =====================================================

    st.markdown(
        "##### عدد الموظفين المستحقين لكل فئة وقيمة البدل حسب الجهات الحكومية"
    )


    nature_display = (
        nature_analysis.copy()
    )


    nature_display[
        "قيمة البدل"
    ] = nature_display[
        "قيمة البدل"
    ].map(
        lambda x: f"{x:,.2f}"
    )


    nature_display[
        "إجمالي قيمة البدل"
    ] = nature_display[
        "إجمالي قيمة البدل"
    ].map(
        lambda x: f"{x:,.2f}"
    )


    st.dataframe(
        nature_display,
        use_container_width=True,
        hide_index=True
    )


# =========================================================
# بيانات العقود المختارة
# =========================================================

st.markdown("---")


st.markdown(
    '<div class="section-title">'
    'بيانات العقود'
    '</div>',
    unsafe_allow_html=True
)


st.caption(
    f"يتم عرض {len(filtered_df):,} سجل من نوع «{selected_contract}»"
)


st.dataframe(
    filtered_df,
    use_container_width=True,
    hide_index=True
)


# =========================================================
# إنشاء ملف Excel
# =========================================================

st.markdown("---")


excel_file = create_excel(

    filtered_df=filtered_df,

    contract_type=selected_contract,

    total_contracts=total_contracts,

    total_cost=total_cost,

    total_basic=total_basic,

    total_supplementary=total_supplementary,

    top_entity=top_entity,

    top_category=top_category,

    category_counts=category_counts,

    entity_counts=entity_counts,

    nature_analysis=nature_analysis,

    total_nature_all=total_nature_all,

    total_nature_eligible_all=total_nature_eligible_all

)


# =========================================================
# تحميل التقرير
# =========================================================

st.download_button(

    label="⬇️ تحميل تقرير Excel",

    data=excel_file,

    file_name=(
        f"تقرير_{selected_contract}.xlsx"
    ),

    mime=(
        "application/"
        "vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),

    use_container_width=True

)
